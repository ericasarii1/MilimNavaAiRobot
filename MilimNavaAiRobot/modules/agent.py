# ══════════════════════════════════════════════════════════════════
#   AGENT — otak tingkat lanjut:
#   1. AGENTIC LOOP : AI memutuskan sendiri tools yang dipakai
#                     (search / calc / wiki / cuaca / link) sampai 3 langkah
#   2. DOC READER   : PDF / DOCX / XLSX / TXT dibaca & diinject
#   3. MOOD-AWARE   : deteksi emosi user → sesuaikan tone jawaban
# ══════════════════════════════════════════════════════════════════

import re
import io
import json
import logging

import aiohttp

log = logging.getLogger("milim.agent")

# ════════════════════ 1. AGENTIC LOOP ════════════════════
MAX_STEPS = 3

TOOL_MANIFEST = (
    "\n\nTOOLS YANG BISA KAMU PAKAI — jawab PAKAI JSON MURNI satu baris "
    "(tanpa teks lain) untuk memakai tool:\n"
    '{"tool":"web_search","query":"..."}  → cari berita/info terkini\n'
    '{"tool":"wiki","query":"..."}       → fakta ensiklopedia\n'
    '{"tool":"weather","place":"..."}    → cuaca kota\n'
    '{"tool":"calc","expression":"..."}  → hitung matematika presisi\n'
    '{"tool":"read_url","url":"..."}     → baca isi halaman web\n\n'
    "Atau jawab langsung pertanyaannya (tanpa JSON) kalau tidak butuh tool. "
    "Kamu boleh memakai tool hingga 3 kali berurutan sebelum jawab final.")


async def run_tool(tool: str, **kw) -> str:
    """Eksekusi satu tool → hasil teks."""
    from MilimNavaAiRobot.modules import web_search as WS
    from MilimNavaAiRobot.modules import smart_tools as ST
    try:
        if tool == "web_search":
            return (await WS.web_search(kw.get("query", ""), limit=4)) or \
                "(tidak ada hasil)"
        if tool == "wiki":
            return (await ST.read_url(
                "https://id.wikipedia.org/wiki/" +
                kw.get("query", "").replace(" ", "_")))[:3000] or "(kosong)"
        if tool == "weather":
            from MilimNavaAiRobot.modules.web_search import _weather
            return (await _weather(kw.get("place", "Jakarta"))) or "(kosong)"
        if tool == "calc":
            return ST.try_calculate(kw.get("expression", "")) or "(bukan ekspresi valid)"
        if tool == "read_url":
            return (await ST.read_url(kw.get("url", "")))[:3000] or "(gagal baca)"
    except Exception as e:
        log.debug(f"tool {tool} err: {e}")
        return f"(tool error: {e})"
    return "(tool tidak dikenal)"


JSON_TOOL_RE = re.compile(r'\{\s*"tool"\s*:\s*"([a-z_]+)"[^}]*\}')


def parse_tool_call(text: str):
    """Return (tool, kwargs) atau None."""
    m = JSON_TOOL_RE.search(text or "")
    if not m:
        return None
    try:
        obj = json.loads(m.group(0))
        tool = obj.pop("tool")
        return tool, obj
    except Exception:
        return None


async def agentic_chat(llm, msgs: list) -> str:
    """Loop: LLM boleh pakai tool hingga MAX_STEPS sebelum jawab final."""
    convo = list(msgs)
    for step in range(MAX_STEPS):
        out = await llm.chat(convo)
        call = parse_tool_call(out)
        if not call:
            return out          # jawaban final
        tool, kw = call
        log.info(f"agent step {step+1}: {tool} {kw}")
        result = await run_tool(tool, **kw)
        convo = convo + [
            {"role": "assistant", "content": out[:300]},
            {"role": "user", "content":
             f"HASIL TOOL {tool}:\n{result[:2500]}\n\n"
             "Lanjut: pakai tool lain (JSON) atau tulis jawaban final."}]
    # habis langkah → paksa jawaban final
    convo.append({"role": "user", "content":
                  "Batas tool habis. Tulis jawaban final sekarang berdasarkan "
                  "semua hasil tool yang sudah didapat."})
    return await llm.chat(convo)


# ════════════════════ 2. DOC READER ════════════════════
MAX_DOC_CHARS = 8000
DOC_MIMES = ("application/pdf",
             "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
             "text/plain", "text/markdown", "text/csv")


def _read_pdf(raw: bytes) -> str:
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(raw))
    out = []
    for i, page in enumerate(r.pages[:30]):
        out.append(f"[hal {i+1}] " + (page.extract_text() or ""))
    return "\n".join(out)


def _read_docx(raw: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(raw))
    return "\n".join(p.text for p in d.paragraphs if p.text.strip())


def _read_xlsx(raw: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True,
                                data_only=True)
    out = []
    for ws in wb.worksheets[:5]:
        out.append(f"[sheet: {ws.title}]")
        for row in list(ws.iter_rows(values_only=True))[:60]:
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


async def read_document(client, msg) -> str:
    """Download & ekstrak isi dokumen → teks, atau '' kalau gagal."""
    d = msg.document
    mime = d.mime_type or ""
    fname = d.file_name or ""
    try:
        path = await client.download_media(d.file_id, file_name="/tmp/milim_doc")
        with open(path, "rb") as f:
            raw = f.read()
        import os
        os.remove(path)
    except Exception as e:
        log.debug(f"doc dl err: {e}")
        return ""
    try:
        if mime == "application/pdf" or fname.lower().endswith(".pdf"):
            text = _read_pdf(raw)
        elif mime.endswith("wordprocessingml.document") or \
                fname.lower().endswith((".docx",)):
            text = _read_docx(raw)
        elif mime.endswith("spreadsheetml.sheet") or \
                fname.lower().endswith((".xlsx",)):
            text = _read_xlsx(raw)
        elif mime.startswith("text/") or fname.lower().endswith(
                (".txt", ".md", ".csv", ".json", ".py", ".log")):
            text = raw.decode("utf-8", errors="ignore")
        else:
            return ""
        text = text.strip()
        if len(text) > MAX_DOC_CHARS:
            text = text[:MAX_DOC_CHARS] + "\n…(dipotong)"
        return text
    except Exception as e:
        log.debug(f"doc parse err: {e}")
        return ""


def is_supported_doc(msg) -> bool:
    d = msg.document
    if not d:
        return False
    mime = d.mime_type or ""
    fn = (d.file_name or "").lower()
    return (mime in DOC_MIMES or mime.startswith("text/") or
            fn.endswith((".pdf", ".docx", ".xlsx", ".txt", ".md",
                         ".csv", ".json", ".py", ".log")))


# ════════════════════ 3. MOOD-AWARE ════════════════════
MOOD_PATTERNS = [
    ("sedih", re.compile(
        r"\b(sedih|nangis|capek banget|lelah|payah|nyerah|hampa|sepi|"
        r"galau|patah hati|ditinggal|stress|stres|depresi|muak|penat|"
        r"gak semangat|nggak semangat|down banget)\b", re.I)),
    ("marah", re.compile(
        r"\b(jengkel|kesel|bete|btg|marah|emosi|benci|muak banget|"
        r"nyebelin|goblok|bodoh banget)\b", re.I)),
    ("senang", re.compile(
        r"\b(yey|yeay|hore|senang banget|bahagia|gembira|lulus|diterima|"
        r"naik gaji|menang|berhasil|akhirnya|seru banget|happy)\b", re.I)),
    ("cemas", re.compile(
        r"\b(takut|cemas|khawatir|gugup|panik|bingung banget|deg-degan|"
        r"nggak enak|gak enak|bakat|besok ujian|wawancara)\b", re.I)),
]

MOOD_TONE = {
    "sedih": ("\n\nMOOD USER TERDETEKSI: sedih/lelah. Sesuaikan nada: lebih "
              "lembut, pelan, empati dulu baru (kalau tepat) hiburan ringan "
              "atau tawaran bantuan. Jangan kebanyakan emoji, jangan bercanda "
              "berlebihan."),
    "marah": ("\n\nMOOD USER TERDETEKSI: kesel/marah. Tenangkan dulu, validasi "
              "perasaannya, jangan ambil pihak secara berlebihan, tawarkan "
              "solusi praktis kalau relevan. Tetap tenang & tidak menggurui."),
    "senang": ("\n\nMOOD USER TERDETEKSI: senang/bahagia. Ikut antusias! Rayakan "
               "bersama dia, boleh lebih banyak emoji & energi tinggi."),
    "cemas": ("\n\nMOOD USER TERDETEKSI: cemas/takut. Dengarkan dulu, "
              "rasionalkan tanpa meremehkan, berikan perspektif menenangkan "
              "dan hal konkret yang bisa dia lakukan."),
}


def detect_mood(text: str) -> str | None:
    """Return mood dominan atau None (netral)."""
    if not text or len(text) > 500:
        return None
    for mood, pat in MOOD_PATTERNS:
        if pat.search(text):
            return mood
    return None


def mood_tone_addon(mood: str | None) -> str:
    return MOOD_TONE.get(mood, "") if mood else ""
